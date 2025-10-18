from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.ext.asyncio import create_async_engine
from sqlalchemy import text
from pydantic import BaseModel
from datetime import datetime, date
from typing import Literal
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from fastapi.responses import StreamingResponse
from io import BytesIO
import os
from dotenv import load_dotenv

load_dotenv()


DATABASE_URL = os.getenv("DATABASE_URL")
if not DATABASE_URL:
    raise ValueError("DATABASE_URL environment is not set")

engine = create_async_engine(DATABASE_URL)



app = FastAPI()

# Add CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],  # Vite dev server
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

categories = Literal["Transport", "Food", "Shopping", "Rent", "Bills", "Groceries", "Entertainment"]

class Expenses(BaseModel):
    category: categories
    amount: float
    notes: str | None = None
    date: str

@app.post("/adding_expense")
async def adding_expenses(expenses: Expenses):
    async with engine.begin() as conn:
        result = await conn.execute(
            text("""
            INSERT INTO expenses (category, amount, notes, date)
            VALUES (:category, :amount, :notes, :date)
            RETURNING id
            """),
            {
                "category": expenses.category,
                "amount": expenses.amount,
                "notes": expenses.notes,
                "date": datetime.strptime(expenses.date, "%d-%m-%Y").date()
            }
        )
        new_id = result.fetchone()[0]
    return {"status": "success", "id": new_id}


@app.delete("/deleting_expense/{expense_id}")
async def delete_rows(expense_id: int):
    async with engine.begin() as conn:
        result = await conn.execute(
            text("""
            DELETE FROM expenses WHERE id = :id  
            """),
            {"id": expense_id}
            
        )
    return {"status": "success"}

@app.put("/updating_expense/{expense_id}")
async def update_rows(expense_id: int, expenses: Expenses):
    async with engine.begin() as conn:
        result = await conn.execute(
            text("""
            UPDATE expenses 
            SET category = :category, 
                amount = :amount, 
                notes = :notes,
                date = :date
            WHERE id = :id
            """),
            {
                "id": expense_id,
                "category": expenses.category,
                "amount": expenses.amount,
                "notes": expenses.notes,
                "date": datetime.strptime(expenses.date, "%d-%m-%Y").date()
                }
        )
    return {"status": "success"}

@app.get("/expenses")
async def get_expenses():
    async with engine.begin() as conn:
        result = await conn.execute(
            text("""
            SELECT id, category, amount, notes, date
            FROM expenses
            ORDER BY date DESC
            """)
        )
        rows = result.fetchall()

        expenses = []
        for row in rows:
            expenses.append({
                "id": row[0],
                "category": row[1],
                "amount": float(row[2]),
                "notes": row[3],
                "date": row[4].strftime("%d-%m-%Y")
            })

    return {"expenses": expenses}

@app.get("/monthly_report/{month}/{year}")
async def generate_report(month: int, year: int):
    async with engine.begin() as conn:
        result = await conn.execute(
            text("""
            SELECT category, amount, notes, date
            FROM expenses
            WHERE EXTRACT(MONTH FROM date) = :month
            AND EXTRACT(YEAR FROM date) = :year
            ORDER BY date ASC
             """),
             {"month": month, "year": year}
        )
        rows = result.fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Expenses {month}-{year}"

    headers = ["Category", "Amount", "Notes", "Date"]
    ws.append(headers)

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    total = 0
    for row in rows:
        ws.append([
            row[0],
            float(row[1]),
            row[2] if row[2] else "",
            row[3].strftime("%d-%m-%Y")
        ])
        total += float(row[1])

    # Add total row
    ws.append([])
    total_row = ws.max_row
    ws.append(["", "TOTAL", total, "", ""])
    ws[f"B{total_row + 1}"].font = Font(bold=True)
    ws[f"C{total_row + 1}"].font = Font(bold=True)

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to BytesIO
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    # Return as downloadable file
    return StreamingResponse(
        excel_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=expenses_{month}_{year}.xlsx"}
    )